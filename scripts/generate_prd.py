"""Generate the v0.4 Word PRD from the reviewed Markdown source (no source mutation)."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PRD_v0.4.md"
TARGET = ROOT / "output" / "doc" / "詠翔專案管理系統_PRD_v0.4.docx"


def font_style(style, size, bold=False, color="253548"):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), "Microsoft JhengHei")


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def table_block(doc, lines):
    parsed = [[plain(x.strip()) for x in s.strip().strip("|").split("|")] for s in lines]
    rows = [r for r in parsed if not all(re.fullmatch(r":?-+:?", c.strip()) for c in r)]
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    width = 6.8 / len(rows[0])
    for row_index, values in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(width)
            cell.text = value
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(5)
                para.paragraph_format.space_before = Pt(5)
                para.paragraph_format.line_spacing = 1.1
                for run in para.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "234561")
                cell._tc.get_or_add_tcPr().append(shading)
            elif row_index % 2 == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F0F4F7")
                cell._tc.get_or_add_tcPr().append(shading)
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    source = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(0.72)
    section.left_margin, section.right_margin = Inches(0.73), Inches(0.73)
    section.header_distance, section.footer_distance = Inches(0.3), Inches(0.3)
    font_style(doc.styles["Normal"], 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11)):
        font_style(doc.styles[name], size, True, "234561")
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(14)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    for name in ("List Bullet", "List Number"):
        font_style(doc.styles[name], 10.5)
        doc.styles[name].paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "詠翔專案管理系統  |  PRODUCT REQUIREMENTS"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("66788B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("v0.4  ·  2026-09-25  |  第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 頁")
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("66788B")
    doc.core_properties.title = "詠翔專案管理系統 PRD v0.4"
    doc.core_properties.subject = "九節點 SOP、V4 日報、期限審批與設計變更"
    doc.core_properties.author = "詠翔專案管理需求整理"
    doc.core_properties.keywords = "PRD, v0.4, Meego, Lark, Zeabur, prototype"
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            table_block(doc, block)
            continue
        if line.startswith("# "):
            doc.add_paragraph(plain(line[2:]), "Title")
        elif line.startswith("### "):
            doc.add_heading(plain(line[4:]), 2)
        elif line.startswith("## "):
            # Keep the concise contents on the opening page, then start the body.
            if line.startswith("## 1. "):
                doc.add_page_break()
            doc.add_heading(plain(line[3:]), 1)
        elif line.startswith("- "):
            doc.add_paragraph(plain(line[2:]), "List Bullet")
        else:
            paragraph = doc.add_paragraph(plain(line))
            if line.startswith("版本："):
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor.from_string("66788B")
        index += 1
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TARGET)
    check = Document(TARGET)
    paragraphs = [p.text for p in check.paragraphs if p.text.strip()]
    tables = [[c.text for r in t.rows for c in r.cells] for t in check.tables]
    all_text = "\n".join(paragraphs + [x for t in tables for x in t])
    checks = ["九大節點", "superseded", "三線", "下一工作日", "PostgreSQL", "Asia/Taipei", "1 營業額", "production", "待審", "API_CONTRACT.md"]
    missing = [value for value in checks if value not in all_text]
    assert not missing, missing
    expected_tables = sum(1 for i, value in enumerate(lines) if value.startswith("|") and (i == 0 or not lines[i - 1].startswith("|")))
    assert len(check.tables) == expected_tables, (len(check.tables), expected_tables)
    assert "�" not in all_text
    print(f"Generated: {TARGET}")
    print(f"Structure verified: {len(paragraphs)} paragraphs, {len(check.tables)} tables, {len(all_text)} characters")
    print("Visual rendering: not verified; LibreOffice and Poppler unavailable in this environment.")


if __name__ == "__main__":
    main()
